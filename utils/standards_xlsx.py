"""Build the Standards directory as a structured Excel workbook.

Sheets:
  Index          — every glossary term that has standards, with its Glossary ID.
  All Standards  — one row per standard (its unique glossary-derived ID + details).
  <per term>     — one sheet per glossary term listing the standards that reference
                   it, each row carrying the standard's unique ID.

IDs: each glossary term has GL-#### ; each standard is owned by one primary term and
gets a unique ID GL-####-S## used everywhere it appears.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

GREEN = '007451'
WHITE = 'FFFFFF'
HEAD_FILL = PatternFill('solid', fgColor=GREEN)
HEAD_FONT = Font(bold=True, color=WHITE)
WRAP = Alignment(vertical='top', wrap_text=True)
THIN = Border(bottom=Side(style='thin', color='DDDDDD'))


def _san_sheet(name, used):
    for ch in '[]:*?/\\':
        name = name.replace(ch, ' ')
    name = (name[:31].strip() or 'Sheet')
    base, k = name, 1
    while name in used:
        suffix = f'_{k}'
        name = (base[:31 - len(suffix)] + suffix)
        k += 1
    used.add(name)
    return name


def _write_header(ws, cols, row=1):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def build_standards_xlsx(path, index_rows, all_rows, sheets, meta):
    wb = Workbook()

    # ---- Index ----
    ws = wb.active
    ws.title = 'Index'
    ws['A1'] = 'AU ECED-FLN Standards Directory'
    ws['A1'].font = Font(bold=True, size=14, color=GREEN)
    ws['A2'] = f"{meta.get('standards', 0)} standards across {len(index_rows)} glossary terms · updated {meta.get('date', '')}"
    ws['A2'].font = Font(color='777777')
    _write_header(ws, ['Glossary ID', 'Glossary Term', '# Standards'], row=4)
    r = 5
    for glid, term, n in index_rows:
        ws.cell(row=r, column=1, value=glid)
        ws.cell(row=r, column=2, value=term)
        ws.cell(row=r, column=3, value=n)
        r += 1
    _widths(ws, [14, 55, 14])

    # ---- All Standards ----
    wsa = wb.create_sheet('All Standards')
    cols = ['Standard ID', 'Title', 'Statement', 'Rationale', 'Category',
            'Primary Glossary Term', 'Related Glossary Terms', 'Source Initiatives']
    _write_header(wsa, cols)
    r = 2
    for row in all_rows:
        vals = [row['sid'], row['title'], row['statement'], row['rationale'], row['category'],
                row['primary_term'], row['related'], row['sources']]
        for j, v in enumerate(vals, start=1):
            cell = wsa.cell(row=r, column=j, value=v)
            cell.alignment = WRAP
            cell.border = THIN
        r += 1
    _widths(wsa, [16, 38, 60, 40, 22, 26, 34, 40])

    # ---- per glossary-term sheets ----
    used = {'Index', 'All Standards'}
    for sh in sheets:
        ws = wb.create_sheet(_san_sheet(f"{sh['glid']} {sh['term']}", used))
        ws['A1'] = sh['term']
        ws['A1'].font = Font(bold=True, size=13, color=GREEN)
        ws['A2'] = f"Glossary ID: {sh['glid']}"
        ws['A2'].font = Font(color='777777')
        if sh.get('aliases'):
            ws['A3'] = 'Also known as: ' + ', '.join(sh['aliases'])
            ws['A3'].font = Font(italic=True, color='777777')
        hrow = 5
        cols2 = ['Standard ID', 'Title', 'Statement', 'Rationale', 'Category',
                 'Primary here?', 'Source Initiatives']
        _write_header(ws, cols2, row=hrow)
        r = hrow + 1
        for row in sh['rows']:
            vals = [row['sid'], row['title'], row['statement'], row['rationale'], row['category'],
                    'Yes' if row['primary'] else 'reference', row['sources']]
            for j, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=j, value=v)
                cell.alignment = WRAP
                cell.border = THIN
            r += 1
        _widths(ws, [16, 38, 60, 40, 22, 14, 40])

    wb.save(path)
    return path
